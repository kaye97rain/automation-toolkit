import pandas as pd
from openpyxl import load_workbook

from core.excel import write_sheets


def test_write_sheets_styles_header(tmp_path):
    path = write_sheets(tmp_path / "out.xlsx", {"Data": pd.DataFrame({"a": [1, 2], "b": ["x", "y"]})})
    ws = load_workbook(path)["Data"]
    assert ws["A1"].value == "a"
    assert ws["A1"].font.bold
    assert ws.freeze_panes == "A2"
