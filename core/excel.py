"""Write pandas DataFrames to nicely formatted Excel workbooks."""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_sheets(path: Path, sheets: dict[str, pd.DataFrame]) -> Path:
    """Write {sheet_name: DataFrame} to one workbook with styled headers,
    frozen header row, auto-filter and auto-sized columns."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            for cell in ws[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(df.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)])
                ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    return path
